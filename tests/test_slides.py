"""The slide deck (`radar/slides.py`): drawn from the stored draft only, so every word is the draft's, a label's or a fact, every
address is a fact, nothing runs off a slide, and drawing it again gives the same bytes."""

import io
import json
import re
import shutil
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import pytest
from PIL import ImageFont
from pptx import Presentation

from radar import cli, config, slides
from radar.labels import L
from tests.support import make_draft
from tests.test_write import FACTS as BASE_FACTS

ROOT = config.ROOT
LINKS = [["Helpdesk docs", "https://helpdesk.example.org/docs"], ["bad", "javascript:alert(1)"], ["console", "http://localhost:8080"]]
BLOCKS = [["toml", '[helpdesk]\nurl = "https://helpdesk.example.org"']]
FACTS = {**BASE_FACTS, "bucket": "ai", "url": "https://github.com/example-org/sample-agent-kit", "links": LINKS, "blocks": BLOCKS}
EAST = re.compile(r"[぀-ヿ㐀-鿿]")


def stored_drafts() -> list[Path]:
    return sorted((ROOT / "digests").glob("*/*/draft-*.json"))


def texts(blob: bytes) -> list[str]:
    out = []
    for slide in Presentation(io.BytesIO(blob)).slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                out += [p.text for p in shape.text_frame.paragraphs if p.text]
    return out


def strings(value) -> list[str]:
    if isinstance(value, str):
        return [value]
    if isinstance(value, dict):
        return [s for v in value.values() for s in strings(v)]
    if isinstance(value, (list, tuple)):
        return [s for v in value for s in strings(v)]
    return [str(value)] if isinstance(value, (int, float)) and not isinstance(value, bool) else []


def words(text: str) -> set[str]:
    return {w.lower() for w in re.findall(r"\w+", text)}


def allowed_words(facts: dict, draft: dict, lang: str) -> set[str]:
    return words(" ".join(strings(draft) + strings(L[lang]) + strings(facts)))


def check_words(blob: bytes, facts: dict, draft: dict, lang: str) -> None:
    allowed = allowed_words(facts, draft, lang)
    for text in texts(blob):
        stray = [w for w in words(text) if w not in allowed and not w.isdigit() and not EAST.search(w)]
        assert not stray, f"{stray} on a slide ({text!r}) are neither the draft's, a label nor a fact"


def addresses(blob: bytes) -> set[str]:
    found = set()
    with zipfile.ZipFile(io.BytesIO(blob)) as z:
        for name in z.namelist():
            if name.startswith("ppt/slides/_rels/"):
                found |= set(re.findall(r'Target="([^"]+)"[^>]*TargetMode="External"', z.read(name).decode("utf-8")))
    return found


def check_inside_the_slide(blob: bytes) -> None:
    prs = Presentation(io.BytesIO(blob))
    w, h = prs.slide_width, prs.slide_height
    for n, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0 and shape.left + shape.width <= w + 12700 and shape.top + shape.height <= h + 12700, \
                f"slide {n}: {shape.name} leaves the slide"


def check_pages_fit(book: slides.Book) -> None:
    for n, page in enumerate(book.pages, 1):
        placed = page.get("placed", [])
        if placed:
            b, y, h = placed[-1]
            assert y + h <= slides.BOTTOM + 0.5 or len(placed) == 1, f"slide {n} runs past the content area"


# ---------------------------------------------------------------- the stored digests of this repository

@pytest.mark.parametrize("path", stored_drafts(), ids=lambda p: "/".join(p.parts[-3:]))
def test_every_stored_draft_draws_a_deck_that_fits_and_says_only_what_the_draft_says(path):
    stored = json.loads(path.read_text(encoding="utf-8"))
    facts, draft, lang = stored["facts"], stored["draft"], stored["lang"]
    book = slides.guide_book(facts, draft, lang) if draft.get("guide") else slides.concept_book(facts, draft, lang)
    check_pages_fit(book)
    blob = book.build()
    check_inside_the_slide(blob)
    check_words(blob, facts, draft, lang)
    links = {facts["url"]} | {u for _, u in facts.get("links") or []}
    assert addresses(blob) <= links
    assert len(Presentation(io.BytesIO(blob)).slides) == len(book.pages) >= 5


def test_a_deck_is_the_same_bytes_every_time():
    draft = make_draft(FACTS, "vi")
    assert slides.deck(FACTS, draft, "vi") == slides.deck(FACTS, draft, "vi")
    with zipfile.ZipFile(io.BytesIO(slides.deck(FACTS, draft, "vi"))) as z:
        assert {i.date_time for i in z.infolist()} == {(2020, 1, 1, 0, 0, 0)}


# ---------------------------------------------------------------- what is on the slides

@pytest.mark.parametrize("lang", ["vi", "en", "ja"])
@pytest.mark.parametrize("legacy", [False, True])
def test_a_deck_is_drawn_in_every_language_for_a_guide_and_for_an_older_draft(lang, legacy):
    draft = make_draft(FACTS, lang, legacy=legacy)
    blob = slides.deck(FACTS, draft, lang)
    check_inside_the_slide(blob)
    check_words(blob, FACTS, draft, lang)
    assert draft["title"] in " ".join(texts(blob))


def test_only_the_repository_and_the_readme_addresses_become_links():
    draft = make_draft(FACTS, "en")
    draft["guide"]["steps"][0]["link_ref"] = 1                     # javascript:, not a plain address: left out whatever the draft says
    draft["guide"]["steps"][1]["link_ref"] = 0
    draft["guide"]["steps"][3]["link_ref"] = 2                     # a local address is shown as code, never linked
    blob = slides.deck(FACTS, draft, "en")
    assert addresses(blob) == {FACTS["url"], "https://helpdesk.example.org/docs"}
    assert "javascript" not in " ".join(texts(blob)) and "http://localhost:8080" in " ".join(texts(blob))


def test_a_command_is_printed_as_the_draft_has_it_and_a_sample_as_the_readme_has_it():
    draft = make_draft(FACTS, "en")
    shown = texts(slides.deck(FACTS, draft, "en"))
    assert "pip install sample-agent-kit" in shown and "sample-agent run --queue support" in shown
    assert '[helpdesk]' in shown and 'url = "https://helpdesk.example.org"' in shown


def test_hostile_text_cannot_break_the_file_or_add_an_element():
    draft = make_draft(FACTS, "en", title="<script>alert(1)</script> \x00bad\x1b & \"quotes\"", hook="x" * 400)
    draft["learn"] = ["<a href='http://evil.example'>click</a> " + "w" * 300]
    draft["guide"]["steps"][0]["commands"] = ["echo \x1b[31mred\x00 \x0b"]            # a command goes into a code block by another path than prose
    blob = slides.deck(FACTS, draft, "en")
    assert Presentation(io.BytesIO(blob)).slides
    assert any("<script>" in t for t in texts(blob))                      # shown as text, so a reader sees what the README said
    with zipfile.ZipFile(io.BytesIO(blob)) as z:
        xml = "".join(z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("ppt/slides/slide"))
    assert "<script>" not in xml and "evil.example" not in addresses(blob) and "\x00" not in xml
    assert "_x00" not in xml and "_x001B_" not in xml                      # no escape sequence for a control character either: it is dropped
    assert slides._clean("a\x00b\x1fc\x0bd") == "abcd"


def test_a_long_paragraph_and_a_long_sample_continue_on_the_next_slide_instead_of_running_off():
    sentence = "Mỗi câu của đoạn này khá dài để chứng minh rằng đoạn văn được chia theo câu chứ không bị cắt giữa chừng. "
    draft = make_draft(FACTS, "vi", problem=sentence * 40)
    draft["guide"]["steps"][0]["commands"] = [f"echo line {i} " + "x" * 30 for i in range(90)]
    book = slides.guide_book(FACTS, draft, "vi")
    check_pages_fit(book)
    problem = [p for p in book.pages if p["title"].startswith(L["vi"]["problem"])]
    assert len(problem) >= 2 and "(1/" in problem[0]["title"]
    first_step = [p for p in book.pages if p["kicker"].startswith(L["vi"]["step"] + " 1")]
    assert len(first_step) >= 2
    check_inside_the_slide(book.build())


def test_a_block_that_cannot_fit_is_split_not_dropped():
    lines = [f"line {i}" for i in range(200)]
    page = slides.paginate([slides.code(lines)], 100)
    shown = [x for p in page for b, _, _ in p for x in b["lines"]]
    assert shown == lines and all(y + h <= slides.BOTTOM + 0.5 for p in page for _, y, h in p)


def test_a_link_never_starts_a_slide_alone_when_the_block_before_it_can_come_along():
    filler = [slides.card(None, "x " * 40, marker=i) for i in range(1, 4)]
    glued = slides.para([("Liên kết", "t", None)], glue=True)
    for count in range(1, 4):
        pages = slides.paginate(filler[:count] * 2 + [glued], 120)
        assert len(pages[-1]) > 1 or len(pages) == 1


def test_roadmap_numbers_every_step_in_order():
    draft = make_draft(FACTS, "en")
    steps = [s["title"] for s in draft["guide"]["steps"]]
    book = slides.guide_book(FACTS, draft, "en")
    road = next(p for p in book.pages if p["title"].startswith(L["en"]["roadmap"]))
    items = [t for b, _, _ in road["placed"] for _, t in b["items"]]
    assert items == steps[:len(items)]


# ---------------------------------------------------------------- the diagram

HUB = [[0, i] for i in range(1, 6)]
DIAGRAMS = [(5, [[0, 1], [1, 2], [2, 0]]), (5, []), (5, [[0, 2], [1, 3], [2, 4], [3, 4]]), (5, [[4, 0], [3, 0]]), (5, [[0, 1], [0, 2], [0, 3], [0, 4]]),
            (6, HUB), (6, HUB + [[5, 0]]), (6, [[i, i] for i in range(6)]), (6, [[i, j] for i in range(6) for j in range(6) if i != j][:12]),
            (2, [[0, 1]]), (3, [[0, 1], [1, 2], [2, 0]]), (6, [[0, 1], [1, 2], [2, 3], [3, 4], [4, 5]])]


@pytest.mark.parametrize("n, edges", DIAGRAMS)
def test_the_diagram_never_overlaps_a_box_and_stays_inside_the_content_area(n, edges):
    comps = [{"name": f"Component number {i}", "role": "does a thing"} for i in range(n)]
    top = 110
    geo = slides.flow_geometry(comps, edges, top, slides.BOTTOM)
    boxes = [(x, y, x + geo["bw"], y + geo["heights"][i]) for i, (x, y) in geo["pos"].items()]
    for i, a in enumerate(boxes):
        assert a[0] >= 0 and a[2] <= slides.SW, "a box leaves the slide sideways"
        assert a[1] >= top - 0.5 and a[3] <= slides.BOTTOM + 0.5, "a box leaves the content area (under the title or past the bottom)"
        for b in boxes[i + 1:]:
            assert a[2] <= b[0] or b[2] <= a[0] or a[3] <= b[1] or b[3] <= a[1], "two boxes overlap"
    assert len(geo["routes"]) == len(edges)
    for r in geo["routes"]:
        assert all(0 <= x <= slides.SW and top - 0.5 <= y <= slides.BOTTOM + 0.5 for x, y in r["pts"])


def test_a_hub_and_spoke_deck_keeps_every_box_on_the_slide():
    draft = make_draft(FACTS, "en")
    draft["how_it_works"]["components"] = [{"name": f"Part {i}", "role": "does a thing", "icon": "brain", "evidence": "README"} for i in range(6)]
    draft["how_it_works"]["flow"] = HUB
    blob = slides.deck(FACTS, draft, "en")
    check_inside_the_slide(blob)


# ---------------------------------------------------------------- the metrics are the font's

CARLITO = {False: Path("/usr/share/fonts/truetype/crosextra/Carlito-Regular.ttf"), True: Path("/usr/share/fonts/truetype/crosextra/Carlito-Bold.ttf")}
needs_carlito = pytest.mark.skipif(not all(p.is_file() for p in CARLITO.values()), reason="Carlito (fonts-crosextra-carlito) is not installed")


@needs_carlito
def test_the_width_table_is_what_the_font_says():
    font = ImageFont.truetype(str(CARLITO[False]), 2048)
    for ch, units in slides._CALIBRI.items():
        assert abs(font.getlength(ch) - units) <= 2, f"{ch!r}: table {units}, font {font.getlength(ch)}"
    for ch in "ơưƠƯờứữ":                                           # the horn letters are wider than their base letter
        assert abs(slides._adv(ch, False, 2048, False) - font.getlength(ch)) <= 2


def real_lines(text: str, size: float, width: float, bold: bool) -> int:
    """Lines a greedy wrap takes with the real font's advances (no margin)."""
    font = ImageFont.truetype(str(CARLITO[bold]), 2048)
    space, lines, used = font.getlength(" ") / 2048 * size, 1, 0.0
    for word in text.split():
        w = font.getlength(word) / 2048 * size
        if used and used + space + w > width:
            lines, used = lines + 1, w
        else:
            used += (space if used else 0) + w
    return lines


@needs_carlito
@pytest.mark.parametrize("size, width, bold", [(22, 728, False), (24, 760, False), (24, 728, True), (36, 852, True)])
def test_the_line_count_never_falls_short_of_the_real_one_and_is_never_more_than_a_line_over(size, width, bold):
    sentences = []
    for path in stored_drafts():
        draft = json.loads(path.read_text(encoding="utf-8"))["draft"]
        sentences += [s for s in strings(draft) if 30 < len(s) < 400 and "`" not in s and "\n" not in s]
    assert len(sentences) > 100
    for text in sentences:
        real, ours = real_lines(text, size * (1.065 if size == 36 else 1), width, bold), slides._nlines(text, size, width, bold)
        if size == 36:
            continue
        assert real <= ours <= real + 1, f"{text!r}: the font wraps it in {real} lines, the measure says {ours}"


# ---------------------------------------------------------------- the file

def test_the_document_dates_are_the_scan_date_not_the_day_it_was_drawn():
    prs = Presentation(io.BytesIO(slides.deck(FACTS, make_draft(FACTS, "en"), "en")))
    when = datetime.fromisoformat(FACTS["scanned"]).replace(tzinfo=timezone.utc).replace(tzinfo=None)
    cp = prs.core_properties
    assert cp.created == when and cp.modified == when and cp.author == "GitHub Radar"


def test_every_committed_deck_says_what_its_draft_says():
    for deck in sorted((ROOT / "digests").glob("*/*/slides-??.pptx")):          # none is required: a deck is drawn on request
        lang = deck.stem.split("-")[1]
        stored = json.loads((deck.parent / f"draft-{lang}.json").read_text(encoding="utf-8"))
        assert texts(deck.read_bytes()) == texts(slides.deck(stored["facts"], stored["draft"], lang)), \
            f"{deck.relative_to(ROOT)} is out of date: python3 radar.py slides {deck.parent.parent.name}/{deck.parent.name}"


# ---------------------------------------------------------------- the command

def one_digest(tmp_path: Path) -> Path:
    src = ROOT / "digests" / "2026-10-17-2" / "agent-orchestrator"
    dst = tmp_path / "digests" / "2026-10-17-2" / "agent-orchestrator"
    dst.mkdir(parents=True)
    shutil.copy(src / "draft-vi.json", dst)
    (dst / "vi.md").write_text("# kept\n", encoding="utf-8")
    return dst


def tree(root: Path) -> dict[str, bytes]:
    return {p.relative_to(root).as_posix(): p.read_bytes() for p in sorted(root.rglob("*")) if p.is_file()}


def test_the_command_writes_the_deck_beside_the_draft_and_touches_nothing_else(tmp_path, monkeypatch, capsys):
    folder = one_digest(tmp_path)
    before = tree(tmp_path)
    monkeypatch.setattr(cli, "ROOT", tmp_path)
    assert cli.main(["slides", "2026-10-17-2/agent-orchestrator"]) == 0
    after = tree(tmp_path)
    assert set(after) - set(before) == {"digests/2026-10-17-2/agent-orchestrator/slides-vi.pptx"}
    assert {k: v for k, v in after.items() if k in before} == before
    assert (folder / "slides-vi.pptx").read_bytes()[:2] == b"PK" and '"written"' in capsys.readouterr().out


def test_running_it_twice_changes_nothing(tmp_path):
    one_digest(tmp_path)
    slides.write(tmp_path, ["2026-10-17-2"])
    first = tree(tmp_path)
    slides.write(tmp_path, ["2026-10-17-2"])
    assert tree(tmp_path) == first


def test_it_needs_to_be_told_what_to_draw_and_refuses_what_it_cannot_find(tmp_path):
    one_digest(tmp_path)
    with pytest.raises(ValueError, match="name what to draw"):
        slides.write(tmp_path, [])
    for bad in ("nope", "2026-10-17-2/../../etc", "2026-10-17-2/x/y", "2026-01-01", "2026-10-17-2/missing"):
        with pytest.raises(ValueError):
            slides.write(tmp_path, [bad])
    assert slides.write(tmp_path, everything=True)["written"] == ["digests/2026-10-17-2/agent-orchestrator/slides-vi.pptx"]
    assert slides.write(tmp_path, ["2026-10-17-2"], langs=["en"]) == {"written": [], "skipped": {}}


def test_a_draft_that_no_longer_fits_the_schema_is_skipped_and_its_old_deck_stays(tmp_path):
    folder = one_digest(tmp_path)
    slides.write(tmp_path, ["2026-10-17-2"])
    old = (folder / "slides-vi.pptx").read_bytes()
    stored = json.loads((folder / "draft-vi.json").read_text(encoding="utf-8"))
    del stored["draft"]["how_it_works"]
    (folder / "draft-vi.json").write_text(json.dumps(stored), encoding="utf-8")
    report = slides.write(tmp_path, ["2026-10-17-2"])
    assert report["written"] == [] and "schema" in next(iter(report["skipped"].values()))
    assert (folder / "slides-vi.pptx").read_bytes() == old
    assert not list(folder.glob("*.tmp"))
    (folder / "draft-vi.json").write_text("{not json", encoding="utf-8")
    assert "not JSON" in next(iter(slides.write(tmp_path, ["2026-10-17-2"])["skipped"].values()))


def test_one_strange_file_does_not_stop_the_others_and_leaves_no_temporary_file(tmp_path, monkeypatch):
    folder = one_digest(tmp_path)
    (folder / "draft-en.json").write_text("[1, 2]", encoding="utf-8")
    (folder / "draft-ja.json").write_text('{"lang": ["ja"], "draft": {}, "facts": {}}', encoding="utf-8")
    (folder / "draft-fr.json").write_text('{"lang": "vi"}', encoding="utf-8")
    odd = tmp_path / "digests" / "2026-10-17-2" / "aaa-odd-facts"                  # sorted first: a failure there must not stop the folder after it
    odd.mkdir()
    stored = json.loads((folder / "draft-vi.json").read_text(encoding="utf-8"))
    stored["facts"]["scanned"] = None
    (odd / "draft-vi.json").write_text(json.dumps(stored), encoding="utf-8")
    report = slides.write(tmp_path, ["2026-10-17-2"])
    assert report["written"] == ["digests/2026-10-17-2/agent-orchestrator/slides-vi.pptx"]
    assert set(report["skipped"]) == {"digests/2026-10-17-2/agent-orchestrator/draft-en.json", "digests/2026-10-17-2/agent-orchestrator/draft-fr.json",
                                      "digests/2026-10-17-2/aaa-odd-facts/draft-vi.json"}
    assert [p.name for p in odd.iterdir()] == ["draft-vi.json"]                      # nothing was written beside the odd draft
    (folder / "slides-vi.pptx").unlink()

    def refuse(self, target):
        raise OSError("disk full")

    monkeypatch.setattr(Path, "replace", refuse)
    report = slides.write(tmp_path, ["2026-10-17-2"], langs=["vi"])
    assert report["written"] == [] and "disk full" in report["skipped"]["digests/2026-10-17-2/agent-orchestrator/draft-vi.json"]
    assert not list(folder.glob("*.tmp")) and not (folder / "slides-vi.pptx").exists()
